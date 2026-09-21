#!/usr/bin/env python3
"""
Signal Collect — field validation against real captures

Answers one question for each of the 35 fields needed for failure detection:
did we actually see it arrive, from which collector, on which tool, how often,
and what did it look like.

This reads captured data. It does not consult the precedence map. A field
that is declared and never arrives shows up here as not observed, which is
the whole point.

    python3 signal_validate.py                  # scan the usual places
    python3 signal_validate.py <dir> [<dir>...]
    python3 signal_validate.py --xlsx           # write the spreadsheet
    python3 signal_validate.py --missing        # only what never arrived
"""

import importlib.util
import json
import os
import sys
from collections import defaultdict, Counter

HERE = os.path.dirname(os.path.abspath(__file__))
HOME = os.path.expanduser("~")
DEFAULT_DIRS = [os.path.join(HOME, ".signal", "sessions"),
                "signal_sessions", "signal_sessions/cursor", "data", ".",
                os.path.join(HOME, "project-signal-dashboard", "signal_sessions")]

# the 35 fields the failure modes need
FIELDS = [
 "agent_id", "agent_type", "cache_read_tokens", "chars_added", "chars_removed",
 "command", "compaction_reason", "cost_usd", "delegation_depth", "exit_code",
 "file_change_kind", "file_path", "loop_count", "observable_type",
 "output_size", "outside_workspace", "parent_agent_id", "plan_item",
 "prompt_text", "raw_ref", "reasoning_text", "response_text", "sequence_num",
 "session_id", "status", "stderr", "stdout", "stop_reason", "tool_arguments",
 "tool_name", "tool_result", "total_tokens", "turn_count", "turn_id",
 "workspace",
]

# which failure modes need each field, for the report
NEEDED_BY = {
 "prompt_text": "1.1 2.2 2.3 1.4 2.1", "file_path": "1.1 1.3 1.4 2.3 2.5 2.6 3.1 3.2 3.3",
 "command": "1.1 1.3 1.5 2.6 3.1 3.2 3.3", "workspace": "1.1",
 "tool_name": "1.1 1.2 1.3 1.4 2.2 2.6", "tool_arguments": "1.1 1.2 1.3 2.3 2.6",
 "outside_workspace": "1.1", "plan_item": "1.1 2.3",
 "turn_id": "1.1 1.3 1.4 2.1 2.2 2.3 2.5 2.6 3.1 3.2 3.3",
 "session_id": "1.1 1.3 1.4 1.5 2.1 2.3 3.1 3.2",
 "agent_id": "1.2 2.4 2.5", "agent_type": "1.2", "parent_agent_id": "1.2 2.4 2.5",
 "raw_ref": "1.2 2.4", "delegation_depth": "1.2 2.4",
 "tool_result": "1.3 2.4 2.5", "sequence_num": "1.3 2.6",
 "loop_count": "1.3 1.5 2.2", "compaction_reason": "1.4 2.1",
 "cache_read_tokens": "1.4 2.1", "total_tokens": "1.4 1.5 2.1",
 "turn_count": "1.5", "exit_code": "1.5 3.1 3.2 3.3", "cost_usd": "1.5",
 "response_text": "1.5 2.2 2.4 2.5 2.6 3.1 3.2",
 "observable_type": "2.1 3.1", "reasoning_text": "2.2 2.4 2.5 2.6",
 "chars_added": "2.3", "chars_removed": "2.3", "status": "3.1",
 "stop_reason": "3.1", "file_change_kind": "3.2 3.3", "stdout": "3.2 3.3",
 "stderr": "3.3", "output_size": "3.3",
}


def load_extractors():
    for name in ("signal_curate.py",):
        p = os.path.join(HERE, name)
        if os.path.exists(p):
            spec = importlib.util.spec_from_file_location("sc", p)
            m = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(m)
            return m
    sys.exit("signal_curate.py must sit beside this script")


SC = load_extractors()


def scan(dirs):
    """Return observed[(field, tool, collector)] = (count, example)."""
    import glob
    from datetime import datetime, timezone
    observed = defaultdict(lambda: [0, None])
    rows_by = Counter()
    seen_files = set()

    def note(field, tool, coll, value):
        if value in (None, "", [], {}):
            return
        k = (field, tool, coll)
        observed[k][0] += 1
        if observed[k][1] is None:
            s = json.dumps(value, default=str) if isinstance(value, (dict, list)) else str(value)
            observed[k][1] = s[:90]

    def record(coll, fields):
        tool = fields.get("tool", "unknown")
        if tool not in ("claude-code", "cursor"):
            return
        rows_by[(tool, coll)] += 1
        for f in FIELDS:
            note(f, tool, coll, fields.get(f))

    for d in dirs:
        if not os.path.isdir(d):
            continue
        for path in glob.glob(os.path.join(d, "**", "*.json*"), recursive=True):
            rp = os.path.realpath(path)
            if rp in seen_files or rp.endswith(".wire.jsonl"):
                continue
            seen_files.add(rp)
            base = os.path.basename(rp).lower()
            for line in open(rp, errors="replace"):
                line = line.strip()
                if not line:
                    continue
                try:
                    rec = json.loads(line)
                except json.JSONDecodeError:
                    continue

                if isinstance(rec, dict) and isinstance(rec.get("payload"), dict):
                    p = rec["payload"]
                    when = SC.ts(rec.get("received_at")) or datetime.now(timezone.utc)
                    is_cursor = "cursor" in base or "cursor_version" in p
                    f = (SC.extract_cursor_hook(p, when) if is_cursor
                         else SC.extract_claude_hook(p, when))
                    record("local_hooks", f)

                elif isinstance(rec, dict) and "resourceLogs" in rec:
                    for rl in rec["resourceLogs"]:
                        res = {a["key"]: SC.unwrap(a.get("value"))
                               for a in rl.get("resource", {}).get("attributes", [])}
                        for sl in rl.get("scopeLogs", []):
                            for lr in sl.get("logRecords", []):
                                at = {a["key"]: SC.unwrap(a.get("value"))
                                      for a in lr.get("attributes", [])}
                                w = SC.ts(at.get("event.timestamp")) or datetime.now(timezone.utc)
                                record("otel", SC.extract_otel(res, at, w))

                elif isinstance(rec, dict) and "observable_id" in rec:
                    coll = rec.get("collector")
                    if coll in ("cli_wrapper", "mcp_proxy", "fs_watcher"):
                        record(coll, SC.extract_row(rec))

    return observed, rows_by, len(seen_files)


COLLECTORS = ["local_hooks", "otel", "cli_wrapper", "mcp_proxy", "fs_watcher"]
SHORT = {"local_hooks": "hooks", "otel": "otel", "cli_wrapper": "wrap",
         "mcp_proxy": "proxy", "fs_watcher": "watch"}


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    dirs = args or DEFAULT_DIRS
    missing_only = "--missing" in sys.argv
    want_xlsx = "--xlsx" in sys.argv

    observed, rows_by, nfiles = scan(dirs)
    if not rows_by:
        print("No captured data found. Looked in:")
        for d in dirs:
            print("   " + d)
        return 1

    print(f"\nscanned {nfiles} files\n")
    print(f"{'tool':<14} {'collector':<14} {'source rows':>12}")
    print("-" * 42)
    for (tool, coll), n in sorted(rows_by.items(), key=lambda x: -x[1]):
        print(f"{tool:<14} {coll:<14} {n:>12}")

    for tool in ("claude-code", "cursor"):
        present = [c for c in COLLECTORS if rows_by.get((tool, c))]
        if not present:
            continue
        print(f"\n\n{'='*78}\n{tool}\n{'='*78}\n")
        head = "  ".join(f"{SHORT[c]:>7}" for c in present)
        print(f"{'field':<22} {head}   modes      example")
        print("-" * 108)
        got = 0
        for f in FIELDS:
            cells, any_here = [], False
            for c in present:
                n = observed.get((f, tool, c), [0, None])[0]
                cells.append(f"{n:>7}" if n else f"{'·':>7}")
                any_here = any_here or bool(n)
            if missing_only and any_here:
                continue
            got += bool(any_here)
            ex = next((observed[(f, tool, c)][1] for c in present
                       if observed.get((f, tool, c), [0, None])[1]), "")
            flag = " " if any_here else "!"
            print(f"{flag}{f:<21} {'  '.join(cells)}   "
                  f"{NEEDED_BY.get(f,''):<10} {(ex or '')[:38]}")
        if not missing_only:
            print(f"\n  observed: {got} of {len(FIELDS)} fields")
            missing = [f for f in FIELDS
                       if not any(observed.get((f, tool, c), [0])[0] for c in present)]
            if missing:
                print(f"  never arrived ({len(missing)}): {', '.join(missing)}")

    if want_xlsx:
        write_xlsx(observed, rows_by)
    return 0


def write_xlsx(observed, rows_by):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    NAVY = "1E2761"; F = "Calibri"
    HDR = PatternFill("solid", fgColor=NAVY)
    G = PatternFill("solid", fgColor="E2EFDA")
    R = PatternFill("solid", fgColor="F4CCCC")
    THIN = Border(*[Side("thin", color="BBBBBB")] * 4)

    wb = Workbook()
    for tool in ("claude-code", "cursor"):
        present = [c for c in COLLECTORS if rows_by.get((tool, c))]
        if not present:
            continue
        ws = wb.active if tool == "claude-code" else wb.create_sheet()
        ws.title = tool
        ws.append(["Field", "Needed by modes"] +
                  [SHORT[c] for c in present] + ["Observed?", "Example value"])
        for c in range(1, len(present) + 5):
            x = ws.cell(row=1, column=c)
            x.font = Font(name=F, size=11, bold=True, color="FFFFFF"); x.fill = HDR
            x.alignment = Alignment(vertical="center", wrap_text=True); x.border = THIN
        for f in FIELDS:
            counts = [observed.get((f, tool, c), [0, None])[0] for c in present]
            any_here = any(counts)
            ex = next((observed[(f, tool, c)][1] for c in present
                       if observed.get((f, tool, c), [0, None])[1]), "")
            ws.append([f, NEEDED_BY.get(f, "")] + [n or "" for n in counts] +
                      ["yes" if any_here else "NEVER", ex or ""])
            r = ws.max_row
            for c in range(1, len(present) + 5):
                x = ws.cell(row=r, column=c); x.font = Font(name=F, size=10)
                x.border = THIN; x.alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row=r, column=1).font = Font(name=F, size=10, bold=True)
            v = ws.cell(row=r, column=len(present) + 3)
            v.fill = G if any_here else R
            v.font = Font(name=F, size=10, bold=True)
        widths = [24, 22] + [8] * len(present) + [12, 46]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = w
        ws.freeze_panes = "A2"
    wb.save(os.path.join(HERE, "field-validation.xlsx"))
    print("\nwritten: field-validation.xlsx")


if __name__ == "__main__":
    sys.exit(main())
